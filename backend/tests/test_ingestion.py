"""Tests for Document Parsing & Ingestion Pipeline (Feature 4).

Covers:
    - POST /api/v1/documents/upload: uploads a sample PDF and confirms status='pending'
    - Confirms document and ACL records exist in PostgreSQL
    - Confirms DocumentVersion and initial DocumentChunks exist with embedding=None
    - POST /api/v1/documents/upload: uploads Markdown document
    - POST /api/v1/documents/upload: uploads DOCX document
    - Edge cases: empty file, invalid roles
"""

import io
import uuid
import pytest
from docx import Document as DocxDocument
from httpx import ASGITransport, AsyncClient
from reportlab.pdfgen import canvas
from sqlalchemy.orm import selectinload

from app.db.session import SyncSessionLocal
from app.main import app
from app.models.document import Document, DocumentACL, DocumentChunk, DocumentVersion


BASE = "/api/v1"


def create_sample_pdf_bytes(title: str = "Company Security Guidelines") -> bytes:
    """Create a minimal valid PDF in memory using ReportLab."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    c.drawString(100, 750, title)
    c.drawString(100, 720, "All employees must enable multi-factor authentication.")
    c.drawString(100, 690, "Security incidents must be reported within 24 hours.")
    c.save()
    return buffer.getvalue()


def create_sample_docx_bytes(title: str = "Engineering Standards") -> bytes:
    """Create a minimal valid DOCX in memory using python-docx."""
    doc = DocxDocument()
    doc.add_heading(title, 0)
    doc.add_paragraph("All code must be reviewed by at least one peer.")
    doc.add_paragraph("Unit test coverage must exceed 80 percent.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Test Cases
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_upload_pdf_lands_in_database_with_status_pending():
    """Upload a sample PDF file and verify it lands in PostgreSQL with status 'pending'."""
    pdf_bytes = create_sample_pdf_bytes("InsideHQ Remote Work Policy")
    files = {
        "file": ("remote_work_policy.pdf", pdf_bytes, "application/pdf")
    }
    data = {
        "title": "Remote Work Policy 2026",
        "department": "Human Resources",
        "allowed_roles": "hr,admin,employee",
    }

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.post(f"{BASE}/documents/upload", files=files, data=data)

    assert resp.status_code == 201, f"Upload failed: {resp.text}"
    body = resp.json()

    assert body["title"] == "Remote Work Policy 2026"
    assert body["source_type"] == "pdf"
    assert body["status"] == "pending"
    assert body["department"] == "Human Resources"
    assert set(body["allowed_roles"]) == {"hr", "admin", "employee"}
    assert body["chunks_count"] > 0
    doc_id = uuid.UUID(body["document_id"])

    # Verify directly in PostgreSQL
    session = SyncSessionLocal()
    try:
        doc = (
            session.query(Document)
            .options(
                selectinload(Document.acls),
                selectinload(Document.versions),
                selectinload(Document.chunks),
            )
            .filter_by(id=doc_id)
            .first()
        )
        assert doc is not None, "Document was not persisted in PostgreSQL"
        assert doc.status == "pending"
        assert doc.source_type == "pdf"
        assert doc.title == "Remote Work Policy 2026"

        # Check ACL entries
        acl_roles = {acl.allowed_role for acl in doc.acls if acl.allowed_role}
        assert acl_roles == {"hr", "admin", "employee"}

        # Check version
        assert len(doc.versions) == 1
        assert doc.versions[0].version_number == 1
        assert doc.versions[0].is_current is True

        # Check chunks and confirm embedding is None (not generated yet)
        assert len(doc.chunks) > 0
        for chunk in doc.chunks:
            assert chunk.embedding is None
            assert chunk.chunk_metadata["department"] == "Human Resources"
            assert "Remote Work Policy" in chunk.chunk_text or "multi-factor" in chunk.chunk_text or len(chunk.chunk_text) > 0
    finally:
        session.close()


@pytest.mark.asyncio
async def test_upload_markdown_document():
    """Upload a Markdown file and verify parsing, status='pending', and ACL entries."""
    md_content = b"# Architecture Overview\n\nInsideHQ uses FastAPI, PostgreSQL pgvector, and React.\n\nMicroservices communicate via secure REST APIs."
    files = {
        "file": ("architecture.md", md_content, "text/markdown")
    }
    data = {
        "title": "System Architecture Overview",
        "department": "Engineering",
        "allowed_roles": '["engineering", "admin"]',
    }

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.post(f"{BASE}/documents/upload", files=files, data=data)

    assert resp.status_code == 201
    body = resp.json()
    assert body["status"] == "pending"
    assert body["source_type"] == "markdown"
    assert set(body["allowed_roles"]) == {"engineering", "admin"}

    session = SyncSessionLocal()
    try:
        doc = session.query(Document).filter_by(id=uuid.UUID(body["document_id"])).first()
        assert doc is not None
        assert doc.status == "pending"
    finally:
        session.close()


@pytest.mark.asyncio
async def test_upload_docx_document():
    """Upload a Word DOCX document and verify parsing into text blocks."""
    docx_bytes = create_sample_docx_bytes("Financial Compliance Rules")
    files = {
        "file": ("compliance.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    }
    data = {
        "title": "Financial Compliance Standards",
        "department": "Finance",
        "allowed_roles": "finance,admin",
    }

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.post(f"{BASE}/documents/upload", files=files, data=data)

    assert resp.status_code == 201
    body = resp.json()
    assert body["status"] == "pending"
    assert body["source_type"] == "docx"
    assert set(body["allowed_roles"]) == {"finance", "admin"}
    assert body["chunks_count"] > 0


@pytest.mark.asyncio
async def test_upload_empty_file_rejected():
    """Empty file upload must return 400 Bad Request."""
    files = {
        "file": ("empty.pdf", b"", "application/pdf")
    }
    data = {
        "title": "Empty Doc",
        "department": "General",
        "allowed_roles": "employee",
    }

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.post(f"{BASE}/documents/upload", files=files, data=data)

    assert resp.status_code == 400
    assert "empty" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_upload_missing_roles_rejected():
    """Upload with empty allowed_roles must return 400 Bad Request."""
    files = {
        "file": ("test.md", b"# Test", "text/markdown")
    }
    data = {
        "title": "Test Doc",
        "department": "General",
        "allowed_roles": "",
    }

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.post(f"{BASE}/documents/upload", files=files, data=data)

    assert resp.status_code in (400, 422)
